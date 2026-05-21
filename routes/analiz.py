# -*- coding: utf-8 -*-
import io
from datetime import datetime
import psycopg2.extras
from flask import Blueprint, jsonify, request, render_template, send_file, make_response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from utils.db import get_db
from utils.query_helpers import login_required, cinsiyet_filter, sezon_filter, donem_filter, ana_grup_filter, marka_filter

analiz_bp = Blueprint('analiz', __name__)


@analiz_bp.route('/')
@login_required
def index():
    resp = make_response(render_template('index.html'))
    resp.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate'
    resp.headers['Pragma'] = 'no-cache'
    return resp


@analiz_bp.route('/api/filters')
@login_required
def filters():
    conn = get_db()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute("SELECT DISTINCT cinsiyet FROM urun_analiz WHERE cinsiyet != 'nan' ORDER BY cinsiyet")
    cinsiyetler = [r['cinsiyet'] for r in cursor.fetchall()]
    cursor.execute("SELECT DISTINCT ana_grup FROM urun_analiz WHERE ana_grup != 'nan' ORDER BY ana_grup")
    gruplar = [r['ana_grup'] for r in cursor.fetchall()]
    cursor.execute("SELECT DISTINCT alt_kategori FROM urun_analiz WHERE alt_kategori != 'nan' ORDER BY alt_kategori")
    kategoriler = [r['alt_kategori'] for r in cursor.fetchall()]
    cursor.execute("SELECT DISTINCT hafta_no FROM urun_analiz WHERE hafta_no IS NOT NULL ORDER BY hafta_no")
    haftalar = [r['hafta_no'] for r in cursor.fetchall()]
    cursor.execute("SELECT DISTINCT sezon FROM urun_analiz WHERE sezon IS NOT NULL ORDER BY sezon")
    sezonlar = [r['sezon'] for r in cursor.fetchall()]
    cursor.execute("SELECT marka FROM (SELECT DISTINCT marka FROM urun_analiz WHERE marka IS NOT NULL AND marka != 'nan') sub ORDER BY LOWER(marka)")
    markalar = [r['marka'] for r in cursor.fetchall()]
    conn.close()
    return jsonify({
        "cinsiyetler": cinsiyetler,
        "ana_gruplar": gruplar,
        "alt_kategoriler": kategoriler,
        "haftalar": haftalar,
        "aylar": [],
        "sezonlar": sezonlar,
        "markalar": markalar
    })


@analiz_bp.route('/api/alt_kategoriler')
@login_required
def alt_kategoriler():
    conn = get_db()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    ana_grup = request.args.get('ana_grup', '')
    if ana_grup:
        ag_list = [a.strip() for a in ana_grup.split(',') if a.strip()]
        if len(ag_list) == 1:
            cursor.execute(
                "SELECT DISTINCT alt_kategori FROM urun_analiz WHERE ana_grup = %s AND alt_kategori != 'nan' ORDER BY alt_kategori",
                (ag_list[0],)
            )
        else:
            placeholders = ','.join(['%s'] * len(ag_list))
            cursor.execute(
                f"SELECT DISTINCT alt_kategori FROM urun_analiz WHERE ana_grup IN ({placeholders}) AND alt_kategori != 'nan' ORDER BY alt_kategori",
                ag_list
            )
    else:
        cursor.execute("SELECT DISTINCT alt_kategori FROM urun_analiz WHERE alt_kategori != 'nan' ORDER BY alt_kategori")
    kategoriler = [r['alt_kategori'] for r in cursor.fetchall()]
    conn.close()
    return jsonify(kategoriler)


@analiz_bp.route('/api/bestseller')
@login_required
def bestseller():
    conn = get_db()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cinsiyet = request.args.get('cinsiyet', '')
    ana_grup = request.args.get('ana_grup', '')
    alt_kategori = request.args.get('alt_kategori', '')
    marka = request.args.get('marka', '')
    donem_tip = request.args.get('donem_tip', 'hafta')
    donem_deger = request.args.get('donem_deger', '1')
    limit = int(request.args.get('limit', 10))
    sezon = request.args.get('sezon', '')
    w_gmroi = float(request.args.get('w_gmroi', 0.40))
    w_st = float(request.args.get('w_st', 0.35))
    w_cover = float(request.args.get('w_cover', 0.25))
    min_st = float(request.args.get('min_st', 0.55))
    max_cover = float(request.args.get('max_cover', 12.0))
    where, params = [], []
    cinsiyet_filter(where, params, cinsiyet)
    ana_grup_filter(where, params, ana_grup)
    marka_filter(where, params, marka)
    if alt_kategori:
        alt_list = [a.strip() for a in alt_kategori.split(',') if a.strip()]
        if len(alt_list) == 1:
            where.append("alt_kategori = %s")
            params.append(alt_list[0])
        elif len(alt_list) > 1:
            where.append("alt_kategori IN (" + ','.join(['%s'] * len(alt_list)) + ")")
            params.extend(alt_list)
    sort_by = request.args.get('sort_by', 'skor')
    sezon_filter(where, params, sezon)
    if sort_by == 'skor':
        where.append("sell_through >= %s")
        params.append(min_st)
        where.append("periyot_cover <= %s")
        params.append(max_cover)
    where_sql = "WHERE " + " AND ".join(where) if where else ""
    bs_sort_map = {
        'st':    ('sell_through',        'DESC'),
        'gmroi': ('COALESCE(gmroi,0)',   'DESC'),
        'cover': ('periyot_cover',       'ASC'),
        'ciro':  ('ciro',                'DESC'),
        'satis': ('toplam_satis_miktar', 'DESC'),
    }
    if sort_by == 'skor':
        score_expr = f"""ROUND((
            {w_gmroi}*COALESCE(NULLIF(gmroi,0)/NULLIF(MAX(NULLIF(gmroi,0)) OVER(),0), 0) +
            {w_st}*COALESCE(sell_through, 0) +
            {w_cover}*(1 - LEAST(periyot_cover/19.0, 1))
        )::numeric, 4) AS bs_skoru"""
        order_clause = "ORDER BY bs_skoru DESC"
    else:
        score_expr = "0 AS bs_skoru"
        col, direction = bs_sort_map.get(sort_by, ('sell_through', 'DESC'))
        order_clause = f"ORDER BY {col} {direction} NULLS LAST"
    query = f"""
    SELECT stok_kodu, stok_kodu_aciklama, marka, ana_grup, alt_kategori,
        cinsiyet, psf, regule_psf, mu, gorsel_url, sezon,
        toplam_satis_miktar AS toplam_satis,
        ROUND(ciro::numeric, 0) AS toplam_ciro,
        ROUND(toplam_kar::numeric, 0) AS toplam_kar,
        indirim_orani AS ort_indirim,
        sell_through,
        periyot_cover,
        toplam_dss_miktar AS dss_miktar,
        ROUND(COALESCE(gmroi, 0)::numeric, 2) AS ort_gmroi,
        sell_through AS avg_st,
        ROUND(periyot_cover::numeric, 1) AS avg_cover,
        {score_expr}
    FROM urun_analiz {where_sql}
    {order_clause}
    LIMIT %s"""
    params.append(limit)
    cursor.execute(query, params)
    sonuclar = cursor.fetchall()
    conn.close()
    return jsonify(list(sonuclar))


@analiz_bp.route('/api/worstseller')
@login_required
def worstseller():
    conn = get_db()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cinsiyet = request.args.get('cinsiyet', '')
    ana_grup = request.args.get('ana_grup', '')
    alt_kategori = request.args.get('alt_kategori', '')
    marka = request.args.get('marka', '')
    donem_tip = request.args.get('donem_tip', 'hafta')
    donem_deger = request.args.get('donem_deger', '1')
    limit = int(request.args.get('limit', 10))
    sezon = request.args.get('sezon', '')
    w_gmroi = float(request.args.get('w_gmroi', 0.20))
    w_st = float(request.args.get('w_st', 0.25))
    w_cover = float(request.args.get('w_cover', 0.55))
    where, params = [], []
    cinsiyet_filter(where, params, cinsiyet)
    ana_grup_filter(where, params, ana_grup)
    marka_filter(where, params, marka)
    if alt_kategori:
        alt_list = [a.strip() for a in alt_kategori.split(',') if a.strip()]
        if len(alt_list) == 1:
            where.append("alt_kategori = %s")
            params.append(alt_list[0])
        elif len(alt_list) > 1:
            where.append("alt_kategori IN (" + ','.join(['%s'] * len(alt_list)) + ")")
            params.extend(alt_list)
    sort_by = request.args.get('sort_by', 'skor')
    sezon_filter(where, params, sezon)
    where_sql = "WHERE " + " AND ".join(where) if where else ""
    ws_sort_map = {
        'st':    ('AVG(sell_through)',        'ASC'),
        'gmroi': ('AVG(NULLIF(gmroi,0))',     'ASC'),
        'cover': ('MAX(periyot_cover)',        'DESC'),
        'ciro':  ('SUM(ciro)',                'ASC'),
        'satis': ('SUM(toplam_satis_miktar)', 'ASC'),
    }
    if sort_by == 'skor':
        score_expr = f"""ROUND((
                {w_gmroi}*(1 - AVG(NULLIF(gmroi,0))/NULLIF(MAX(AVG(NULLIF(gmroi,0))) OVER(),0)) +
                {w_st}*(1 - COALESCE(MAX(sell_through),0)) +
                {w_cover}*LEAST(MAX(periyot_cover)/19.0, 1)
            )::numeric,4) as ws_skoru"""
        order_clause = "ORDER BY ws_skoru DESC, SUM(toplam_satis_miktar) ASC"
    else:
        score_expr = "0 as ws_skoru"
        col, direction = ws_sort_map.get(sort_by, ('AVG(sell_through)', 'ASC'))
        order_clause = f"ORDER BY {col} {direction} NULLS LAST, SUM(toplam_satis_miktar) ASC"
    query = f"""
        SELECT stok_kodu, stok_kodu_aciklama, marka, ana_grup, alt_kategori,
            cinsiyet, psf, regule_psf, mu, gorsel_url, sezon,
            SUM(toplam_satis_miktar) as toplam_satis,
            ROUND(SUM(ciro)::numeric, 0) as toplam_ciro,
            ROUND(SUM(toplam_kar)::numeric, 0) as toplam_kar,
            AVG(indirim_orani) as ort_indirim,
            MAX(sell_through) as sell_through,
            MAX(periyot_cover) as periyot_cover,
            MAX(toplam_dss_miktar) as dss_miktar,
            ROUND(AVG(NULLIF(gmroi,0))::numeric, 2) as ort_gmroi,
            ROUND(AVG(sell_through)::numeric, 4) as avg_st,
            ROUND(AVG(periyot_cover)::numeric, 1) as avg_cover,
            {score_expr}
        FROM urun_analiz {where_sql}
        GROUP BY stok_kodu,stok_kodu_aciklama,marka,ana_grup,alt_kategori,cinsiyet,psf,regule_psf,mu,gorsel_url,sezon
        {order_clause} LIMIT %s"""
    params.append(limit)
    cursor.execute(query, params)
    sonuclar = cursor.fetchall()
    conn.close()
    return jsonify(list(sonuclar))


@analiz_bp.route('/api/export/top10')
@login_required
def export_top10():
    conn = get_db()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

    cinsiyet    = request.args.get('cinsiyet', '')
    ana_grup    = request.args.get('ana_grup', '')
    alt_kategori = request.args.get('alt_kategori', '')
    donem_tip   = request.args.get('donem_tip', 'hafta')
    donem_deger = request.args.get('donem_deger', '1')
    sezon       = request.args.get('sezon', '')
    w_gmroi_bs  = float(request.args.get('w_gmroi', 0.40))
    w_st_bs     = float(request.args.get('w_st', 0.35))
    w_cover_bs  = float(request.args.get('w_cover', 0.25))
    min_st      = float(request.args.get('min_st', 0.55))
    max_cover   = float(request.args.get('max_cover', 12.0))

    def build_where(extra_filters=None):
        where, params = [], []
        cinsiyet_filter(where, params, cinsiyet)
        ana_grup_filter(where, params, ana_grup)
        if alt_kategori:
            alt_list = [a.strip() for a in alt_kategori.split(',') if a.strip()]
            if len(alt_list) == 1:
                where.append("alt_kategori = %s"); params.append(alt_list[0])
            elif len(alt_list) > 1:
                where.append("alt_kategori IN (" + ','.join(['%s']*len(alt_list)) + ")")
                params.extend(alt_list)
        donem_filter(where, params, donem_tip, donem_deger)
        sezon_filter(where, params, sezon)
        if extra_filters:
            for clause, val in extra_filters:
                where.append(clause); params.append(val)
        return ("WHERE " + " AND ".join(where)) if where else "", params

    where_bs, params_bs = build_where([
        ("sell_through >= %s", min_st),
        ("periyot_cover <= %s", max_cover),
    ])
    bs_query = f"""
        SELECT stok_kodu, stok_kodu_aciklama, marka, ana_grup, alt_kategori,
            cinsiyet, psf, mu, sezon,
            toplam_satis_miktar AS toplam_satis,
            ROUND(ciro::numeric, 0) AS toplam_ciro,
            ROUND(toplam_kar::numeric, 0) AS toplam_kar,
            indirim_orani AS ort_indirim,
            sell_through,
            periyot_cover,
            ROUND(COALESCE(gmroi,0)::numeric, 2) AS ort_gmroi,
            ROUND((
                {w_gmroi_bs}*COALESCE(NULLIF(gmroi,0)/NULLIF(MAX(NULLIF(gmroi,0)) OVER(),0),0) +
                {w_st_bs}*COALESCE(sell_through,0) +
                {w_cover_bs}*(1 - LEAST(periyot_cover/19.0,1))
            )::numeric, 4) AS bs_skoru
        FROM urun_analiz {where_bs}
        ORDER BY bs_skoru DESC LIMIT 10"""
    cursor.execute(bs_query, params_bs)
    bs_rows = cursor.fetchall()

    where_ws, params_ws = build_where()
    ws_query = f"""
        SELECT stok_kodu, stok_kodu_aciklama, marka, ana_grup, alt_kategori,
            cinsiyet, psf, mu, sezon,
            SUM(toplam_satis_miktar) as toplam_satis,
            ROUND(SUM(ciro)::numeric, 0) as toplam_ciro,
            ROUND(SUM(toplam_kar)::numeric, 0) as toplam_kar,
            AVG(indirim_orani) as ort_indirim,
            MAX(sell_through) as sell_through,
            MAX(periyot_cover) as periyot_cover,
            ROUND(AVG(NULLIF(gmroi,0))::numeric, 2) as ort_gmroi,
            ROUND((
                0.20*(1 - AVG(NULLIF(gmroi,0))/NULLIF(MAX(AVG(NULLIF(gmroi,0))) OVER(),0)) +
                0.25*(1 - COALESCE(MAX(sell_through),0)) +
                0.55*LEAST(MAX(periyot_cover)/19.0,1)
            )::numeric, 4) as ws_skoru
        FROM urun_analiz {where_ws}
        GROUP BY stok_kodu,stok_kodu_aciklama,marka,ana_grup,alt_kategori,cinsiyet,psf,mu,sezon
        ORDER BY ws_skoru DESC LIMIT 10"""
    cursor.execute(ws_query, params_ws)
    ws_rows = cursor.fetchall()
    conn.close()

    wb = Workbook()

    def _thin():
        s = Side(style='thin', color='D1D5DB')
        return Border(left=s, right=s, top=s, bottom=s)

    def _build_sheet(ws, rows, title, header_hex, score_col_label):
        ws.title = title
        cols = [
            ('Sıra',            8),
            ('Stok Kodu',      14),
            ('Ürün Adı',       36),
            ('Marka',          14),
            ('Alt Kategori',   18),
            ('Cinsiyet',       11),
            ('Sezon',          10),
            ('PSF (₺)',        11),
            ('MU',             10),
            ('Sell Through',   13),
            ('Ciro (₺)',       14),
            ('Kâr (₺)',        13),
            ('İndirim %',      11),
            ('GMROI',          10),
            ('Cover (hafta)',  14),
            (score_col_label,  12),
        ]
        header_fill = PatternFill('solid', fgColor=header_hex)
        header_font = Font(bold=True, color='FFFFFF', size=10)
        center      = Alignment(horizontal='center', vertical='center', wrap_text=False)
        left        = Alignment(horizontal='left',   vertical='center')

        for col_idx, (label, width) in enumerate(cols, 1):
            cell = ws.cell(row=1, column=col_idx, value=label)
            cell.fill   = header_fill
            cell.font   = header_font
            cell.alignment = center
            cell.border = _thin()
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        ws.row_dimensions[1].height = 22

        alt_fill = PatternFill('solid', fgColor='F9FAFB')
        for row_idx, d in enumerate(rows, 2):
            fill = alt_fill if row_idx % 2 == 0 else PatternFill('solid', fgColor='FFFFFF')
            st_val = float(d.get('sell_through') or 0)
            score  = float(d.get('bs_skoru') or d.get('ws_skoru') or 0)
            row_data = [
                row_idx - 1,
                d.get('stok_kodu') or '',
                d.get('stok_kodu_aciklama') or '',
                d.get('marka') or '',
                d.get('alt_kategori') or '',
                d.get('cinsiyet') or '',
                d.get('sezon') or '',
                float(d.get('psf') or 0),
                float(d.get('mu') or 0),
                st_val,
                float(d.get('toplam_ciro') or 0),
                float(d.get('toplam_kar') or 0),
                float(d.get('ort_indirim') or 0) * 100,
                float(d.get('ort_gmroi') or 0),
                float(d.get('periyot_cover') or 0),
                score,
            ]
            num_fmts = [
                None, None, None, None, None, None, None,
                '#,##0.00', '0.0000', '0.00%',
                '#,##0', '#,##0', '0.0"%"', '0.00', '0.0', '0.0000',
            ]
            aligns = [center, left, left, left, left, center, center,
                      center, center, center, center, center, center, center, center, center]
            for col_idx, (val, fmt, aln) in enumerate(zip(row_data, num_fmts, aligns), 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.fill      = fill
                cell.alignment = aln
                cell.border    = _thin()
                cell.font      = Font(size=9)
                if fmt:
                    cell.number_format = fmt
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}1'

    _build_sheet(wb.active,  bs_rows, 'Best Sellers',  '16A34A', 'BS Skoru')
    ws_sheet = wb.create_sheet()
    _build_sheet(ws_sheet,   ws_rows, 'Worst Sellers', 'DC2626', 'WS Skoru')

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    tarih = datetime.now().strftime('%Y%m%d_%H%M')
    filename = f'Top10_{tarih}.xlsx'
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )


@analiz_bp.route('/api/analiz/karsilastirma')
@login_required
def karsilastirma():
    conn = get_db()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cinsiyet = request.args.get('cinsiyet', '')
    ana_grup = request.args.get('ana_grup', '')
    donem_tip = request.args.get('donem_tip', 'hafta')
    donem_deger = request.args.get('donem_deger', '')
    sezon = request.args.get('sezon', '')
    where, params = [], []
    cinsiyet_filter(where, params, cinsiyet)
    ana_grup_filter(where, params, ana_grup)
    donem_filter(where, params, donem_tip, donem_deger)
    sezon_filter(where, params, sezon)
    where_sql = "WHERE " + " AND ".join(where) if where else ""
    cursor.execute(
        f"SELECT SUM(toplam_satis_miktar) as toplam_satis,ROUND(SUM(ciro)::numeric,0) as toplam_ciro,"
        f"ROUND(SUM(toplam_kar)::numeric,0) as toplam_kar,ROUND(AVG(sell_through)::numeric,4) as avg_st "
        f"FROM urun_analiz {where_sql}",
        params
    )
    veri = cursor.fetchone()
    conn.close()
    return jsonify({"bu_donem": dict(veri), "onceki_donem": {}})


@analiz_bp.route('/api/analiz/kategori-dagilim')
@login_required
def kategori_dagilim():
    conn = get_db()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cinsiyet = request.args.get('cinsiyet', '')
    donem_tip = request.args.get('donem_tip', 'hafta')
    donem_deger = request.args.get('donem_deger', '1')
    ana_grup = request.args.get('ana_grup', '')
    sezon = request.args.get('sezon', '')
    where = ["ana_grup != 'nan'", "alt_kategori != 'nan'"]
    params = []
    cinsiyet_filter(where, params, cinsiyet)
    donem_filter(where, params, donem_tip, donem_deger)
    sezon_filter(where, params, sezon)
    ag_list = [a.strip() for a in ana_grup.split(',') if a.strip()] if ana_grup else []
    ana_grup_filter(where, params, ana_grup)
    where_sql = "WHERE " + " AND ".join(where)
    if ag_list:
        cursor.execute(
            f"SELECT alt_kategori as label,SUM(toplam_satis_miktar) as toplam_satis,"
            f"ROUND(SUM(ciro)::numeric,0) as toplam_ciro FROM urun_analiz {where_sql} "
            f"GROUP BY alt_kategori ORDER BY toplam_satis DESC",
            params
        )
    else:
        cursor.execute(
            f"SELECT ana_grup as label,SUM(toplam_satis_miktar) as toplam_satis,"
            f"ROUND(SUM(ciro)::numeric,0) as toplam_ciro FROM urun_analiz {where_sql} "
            f"GROUP BY ana_grup ORDER BY toplam_satis DESC",
            params
        )
    sonuclar = cursor.fetchall()
    conn.close()
    return jsonify(list(sonuclar))



@analiz_bp.route('/api/analiz/en-karli-urunler')
@login_required
def en_karli_urunler():
    conn = get_db()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cinsiyet = request.args.get('cinsiyet', '')
    ana_grup = request.args.get('ana_grup', '')
    donem_tip = request.args.get('donem_tip', 'hafta')
    donem_deger = request.args.get('donem_deger', '1')
    limit = int(request.args.get('limit', 10))
    sezon = request.args.get('sezon', '')
    min_st = float(request.args.get('min_st', 0.55))
    max_cover = float(request.args.get('max_cover', 12.0))
    where, params = [], []
    cinsiyet_filter(where, params, cinsiyet)
    ana_grup_filter(where, params, ana_grup)
    donem_filter(where, params, donem_tip, donem_deger)
    sezon_filter(where, params, sezon)
    where.append("sell_through >= %s")
    params.append(min_st)
    where.append("periyot_cover <= %s")
    params.append(max_cover)
    where_sql = "WHERE " + " AND ".join(where) if where else ""
    cursor.execute(f"""
        SELECT stok_kodu, stok_kodu_aciklama, marka, alt_kategori,
            ROUND(SUM(toplam_kar)::numeric, 0) as toplam_kar,
            ROUND(SUM(ciro)::numeric, 0) as toplam_ciro,
            SUM(toplam_satis_miktar) as toplam_satis,
            ROUND(AVG(mu)::numeric, 2) as avg_mu,
            ROUND(AVG(sell_through)::numeric, 4) as avg_st,
            ROUND(AVG(periyot_cover)::numeric, 1) as periyot_cover,
            ROUND(AVG(NULLIF(gmroi,0))::numeric, 2) as ort_gmroi
        FROM urun_analiz {where_sql}
        GROUP BY stok_kodu, stok_kodu_aciklama, marka, alt_kategori
        ORDER BY SUM(toplam_kar) DESC NULLS LAST
        LIMIT %s""", params + [limit])
    sonuclar = cursor.fetchall()
    conn.close()
    return jsonify(list(sonuclar))


@analiz_bp.route('/api/analiz/en-cok-ciro-urunler')
@login_required
def en_cok_ciro_urunler():
    conn = get_db()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cinsiyet = request.args.get('cinsiyet', '')
    ana_grup = request.args.get('ana_grup', '')
    donem_tip = request.args.get('donem_tip', 'hafta')
    donem_deger = request.args.get('donem_deger', '1')
    limit = int(request.args.get('limit', 8))
    sezon = request.args.get('sezon', '')
    min_st = float(request.args.get('min_st', 0.55))
    max_cover = float(request.args.get('max_cover', 10.0))
    where, params = [], []
    cinsiyet_filter(where, params, cinsiyet)
    ana_grup_filter(where, params, ana_grup)
    donem_filter(where, params, donem_tip, donem_deger)
    sezon_filter(where, params, sezon)
    where.append("sell_through >= %s")
    params.append(min_st)
    where.append("periyot_cover <= %s")
    params.append(max_cover)
    where_sql = "WHERE " + " AND ".join(where) if where else ""
    cursor.execute(f"""
        SELECT stok_kodu, stok_kodu_aciklama, marka, alt_kategori,
            ROUND(SUM(ciro)::numeric, 0) as toplam_ciro,
            ROUND(SUM(toplam_kar)::numeric, 0) as toplam_kar,
            SUM(toplam_satis_miktar) as toplam_satis,
            ROUND(AVG(mu)::numeric, 2) as avg_mu,
            ROUND(AVG(sell_through)::numeric, 4) as avg_st,
            ROUND(AVG(periyot_cover)::numeric, 1) as periyot_cover,
            ROUND(AVG(NULLIF(gmroi,0))::numeric, 2) as ort_gmroi
        FROM urun_analiz {where_sql}
        GROUP BY stok_kodu, stok_kodu_aciklama, marka, alt_kategori
        ORDER BY SUM(ciro) DESC NULLS LAST
        LIMIT %s""", params + [limit])
    sonuclar = cursor.fetchall()
    conn.close()
    return jsonify(list(sonuclar))


@analiz_bp.route('/api/analiz/marka-performans')
@login_required
def marka_performans():
    conn = get_db()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cinsiyet = request.args.get('cinsiyet', '')
    ana_grup = request.args.get('ana_grup', '')
    sezon = request.args.get('sezon', '')
    where, params = ["marka IS NOT NULL", "marka != 'nan'"], []
    cinsiyet_filter(where, params, cinsiyet)
    ana_grup_filter(where, params, ana_grup)
    sezon_filter(where, params, sezon)
    where_sql = "WHERE " + " AND ".join(where)
    # Group by marka only; ciro-weighted averages for ST and GMROI
    cursor.execute(f"""
        SELECT marka,
            ROUND(SUM(ciro)::numeric, 0) AS toplam_ciro,
            ROUND(SUM(toplam_kar)::numeric, 0) AS toplam_kar,
            ROUND(
                (SUM(sell_through * ciro) / NULLIF(SUM(ciro), 0))::numeric, 4
            ) AS avg_st,
            ROUND(
                (SUM(CASE WHEN gmroi > 0 THEN gmroi * ciro ELSE 0 END) /
                 NULLIF(SUM(CASE WHEN gmroi > 0 THEN ciro ELSE 0 END), 0))::numeric, 2
            ) AS avg_gmroi,
            ROUND(
                (SUM(CASE WHEN periyot_cover > 0 THEN periyot_cover * toplam_satis_miktar ELSE 0 END) /
                 NULLIF(SUM(CASE WHEN periyot_cover > 0 THEN toplam_satis_miktar ELSE 0 END), 0))::numeric, 1
            ) AS avg_cover,
            SUM(toplam_satis_miktar) AS toplam_satis
        FROM urun_analiz {where_sql}
        GROUP BY marka
        HAVING SUM(ciro) > 0
        ORDER BY toplam_ciro DESC
    """, params)
    markalar = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return jsonify({"markalar": markalar})


@analiz_bp.route('/api/analiz/en-cok-satan-urunler')
@login_required
def en_cok_satan_urunler():
    conn = get_db()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cinsiyet    = request.args.get('cinsiyet', '')
    ana_grup    = request.args.get('ana_grup', '')
    donem_tip   = request.args.get('donem_tip', 'hafta')
    donem_deger = request.args.get('donem_deger', '')
    sezon       = request.args.get('sezon', '')
    limit       = int(request.args.get('limit', 10))
    where, params = [], []
    cinsiyet_filter(where, params, cinsiyet)
    ana_grup_filter(where, params, ana_grup)
    donem_filter(where, params, donem_tip, donem_deger)
    sezon_filter(where, params, sezon)
    where_sql = "WHERE " + " AND ".join(where) if where else ""
    cursor.execute(f"""
        SELECT stok_kodu, stok_kodu_aciklama, marka,
               SUM(toplam_satis_miktar) AS toplam_satis,
               ROUND(SUM(ciro)::numeric, 0) AS toplam_ciro
        FROM urun_analiz {where_sql}
        GROUP BY stok_kodu, stok_kodu_aciklama, marka
        ORDER BY toplam_satis DESC
        LIMIT %s
    """, params + [limit])
    sonuclar = cursor.fetchall()
    conn.close()
    return jsonify(list(sonuclar))


@analiz_bp.route('/api/analiz/kategori-performans')
@login_required
def kategori_performans():
    conn = get_db()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cinsiyet = request.args.get('cinsiyet', '')
    ana_grup = request.args.get('ana_grup', '')
    donem_tip = request.args.get('donem_tip', 'hafta')
    donem_deger = request.args.get('donem_deger', '1')
    sezon = request.args.get('sezon', '')
    where, params = [], []
    cinsiyet_filter(where, params, cinsiyet)
    ana_grup_filter(where, params, ana_grup)
    donem_filter(where, params, donem_tip, donem_deger)
    sezon_filter(where, params, sezon)
    where_sql = "WHERE " + " AND ".join(where) if where else ""
    cursor.execute(
        f"SELECT alt_kategori,ana_grup,SUM(toplam_satis_miktar) as toplam_satis,"
        f"ROUND(SUM(ciro)::numeric,0) as toplam_ciro,ROUND(SUM(toplam_kar)::numeric,0) as toplam_kar,"
        f"ROUND(AVG(mu)::numeric,2) as avg_mu,ROUND(AVG(sell_through)::numeric,4) as avg_st,"
        f"COUNT(DISTINCT stok_kodu) as urun_sayisi FROM urun_analiz {where_sql} "
        f"GROUP BY alt_kategori,ana_grup ORDER BY toplam_ciro DESC",
        params
    )
    sonuclar = cursor.fetchall()
    conn.close()
    return jsonify(list(sonuclar))
